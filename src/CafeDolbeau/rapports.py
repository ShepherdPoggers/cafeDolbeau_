from calendar import monthrange
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import TransactionCafe


FUSEAU_RAPPORT = ZoneInfo("America/Toronto")
PRECISIONS = (
    ("periode", "Toute la période"),
    ("jour", "Par jour"),
    ("semaine", "Par semaine"),
    ("mois", "Par mois"),
    ("transaction", "Par transaction"),
)
TYPES_RAPPORT = (
    (TransactionCafe.Type.ACHAT, "Achats directs"),
    (TransactionCafe.Type.PREPAYE, "Cafés prépayés"),
    (TransactionCafe.Type.GRATUIT, "Cafés gratuits accordés"),
    (TransactionCafe.Type.UTILISE, "Cafés consommés (prépayés ou gratuits)"),
)


def dernier_jour_mois(date):
    return date.replace(day=monthrange(date.year, date.month)[1])


@dataclass
class QuantitesCategorie:
    code: str
    libelle: str
    enregistrees: int = 0
    annulees: int = 0

    @property
    def net(self):
        return self.enregistrees - self.annulees


@dataclass
class GroupeRapport:
    debut: date
    fin: date
    categories: list
    precision: str

    @property
    def libelle_periode(self):
        if self.precision == "mois":
            return self.debut.strftime("%Y-%m")
        if self.debut == self.fin:
            return self.debut.isoformat()
        return f"{self.debut.isoformat()} au {self.fin.isoformat()}"

    @property
    def total(self):
        return sum(categorie.net for categorie in self.categories)


@dataclass(frozen=True)
class EvenementRapport:
    date: datetime
    transaction_id: int
    type_transaction: str
    nature: str
    client: str
    employe: str
    enregistrees: int
    annulees: int

    @property
    def net(self):
        return self.enregistrees - self.annulees


@dataclass
class Rapport:
    date_debut: date
    date_fin: date
    precision: str
    types: tuple
    genere_le: datetime
    totaux: list
    groupes: list
    evenements: list

    @property
    def total(self):
        return sum(categorie.net for categorie in self.totaux)


def _categories(types):
    categories = [QuantitesCategorie(type_cafe, libelle) for type_cafe, libelle in TYPES_RAPPORT if type_cafe in types]
    return categories, {categorie.code: index for index, categorie in enumerate(categories)}


def _debut_groupe(date, precision):
    if precision == "semaine":
        return date - timedelta(days=date.weekday())
    if precision == "mois":
        return date.replace(day=1)
    return date


def _groupes(debut, fin, precision, categories):
    groupes = []
    courant = debut
    while courant <= fin:
        fin_groupe = courant
        if precision in ("periode", "transaction"):
            fin_groupe = fin
        elif precision == "semaine":
            fin_groupe = courant + timedelta(days=6)
        elif precision == "mois":
            fin_groupe = dernier_jour_mois(courant)
        groupes.append(GroupeRapport(courant, fin_groupe, [
            QuantitesCategorie(categorie.code, categorie.libelle) for categorie in categories
        ], precision))
        courant = fin_groupe + timedelta(days=1)
    return groupes


def generer_rapport(*, date_debut, date_fin, types, precision):
    """Compte les événements à leur date, indépendamment des soldes actuels."""
    debut = datetime.combine(date_debut, time.min, FUSEAU_RAPPORT)
    fin_exclusive = datetime.combine(date_fin + timedelta(days=1), time.min, FUSEAU_RAPPORT)
    categories, colonnes = _categories(types)
    groupes = _groupes(date_debut, date_fin, precision, categories)
    groupes_par_date = {groupe.debut: groupe for groupe in groupes}
    evenements = []
    lignes = TransactionCafe.objects.filter(type_transaction__in=types).filter(
        Q(date_creation__gte=debut, date_creation__lt=fin_exclusive)
        | Q(annulation__date_annulation__gte=debut, annulation__date_annulation__lt=fin_exclusive)
    ).order_by().values(
        "pk", "quantite", "type_transaction", "date_creation", "client__nom_complet",
        "employe__username", "annulation__date_annulation", "annulation__employe__username",
    )
    libelles = dict(TYPES_RAPPORT)
    # Une seule lecture jointe : les événements et leurs traces partagent
    # le même instant de lecture, même pendant une annulation concurrente.
    for ligne in lignes.iterator(chunk_size=1000):
        dates = (
            (ligne["date_creation"], False, ligne["employe__username"]),
            (ligne["annulation__date_annulation"], True, ligne["annulation__employe__username"]),
        )
        for date, annulation, employe in dates:
            if date is None or not debut <= date < fin_exclusive:
                continue
            date_locale = date.astimezone(FUSEAU_RAPPORT)
            cle = date_debut if precision in ("periode", "transaction") else _debut_groupe(date_locale.date(), precision)
            index = colonnes[ligne["type_transaction"]]
            champ = "annulees" if annulation else "enregistrees"
            for quantites in (categories[index], groupes_par_date[cle].categories[index]):
                setattr(quantites, champ, getattr(quantites, champ) + ligne["quantite"])
            if precision == "transaction":
                evenements.append(EvenementRapport(
                    date_locale, ligne["pk"], libelles[ligne["type_transaction"]],
                    "Annulation" if annulation else "Transaction", ligne["client__nom_complet"],
                    employe or "Inconnu", 0 if annulation else ligne["quantite"],
                    ligne["quantite"] if annulation else 0,
                ))
    evenements.sort(key=lambda evenement: (evenement.date.timestamp(), evenement.transaction_id, evenement.nature == "Annulation"))
    return Rapport(date_debut, date_fin, precision, tuple(types),
                   timezone.now().astimezone(FUSEAU_RAPPORT), categories, groupes, evenements)


def _ajouter_ligne(feuille, valeurs):
    feuille.append(valeurs)
    for cellule in feuille[feuille.max_row]:
        if isinstance(cellule.value, str):
            cellule.data_type = "s"
        elif isinstance(cellule.value, datetime):
            cellule.number_format = "dd/mm/yyyy hh:mm:ss"
        elif hasattr(cellule.value, "year"):
            cellule.number_format = "dd/mm/yyyy"
        elif isinstance(cellule.value, int):
            cellule.number_format = "0"


def _mettre_en_forme(feuille, filtre=True):
    feuille.freeze_panes = "A2"
    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="305B96")
    for index, colonne in enumerate(feuille.columns, start=1):
        largeur = max(len(str(cellule.value or "")) for cellule in colonne) + 2
        feuille.column_dimensions[get_column_letter(index)].width = min(50, max(16, largeur))
    if filtre:
        feuille.auto_filter.ref = feuille.dimensions


def _detail_totaux_excel(classeur, rapport):
    feuille = classeur.create_sheet("Détail des totaux")
    _ajouter_ligne(feuille, ["Catégorie", "Enregistrées", "Annulées", "Net"])
    blocs = [
        (groupe.debut if groupe.debut == groupe.fin else groupe.libelle_periode, groupe.categories)
        for groupe in rapport.groupes
    ]
    blocs.append(("TOTAL", rapport.totaux))
    for index, (periode, categories) in enumerate(blocs):
        _ajouter_ligne(feuille, [periode, None, None, None])
        ligne_periode = feuille.max_row
        feuille.merge_cells(start_row=ligne_periode, start_column=1, end_row=ligne_periode, end_column=4)
        for cellule in feuille[ligne_periode]:
            cellule.fill = PatternFill("solid", fgColor="244675")
            cellule.font = Font(bold=True, color="FFFFFF")
            cellule.alignment = Alignment(horizontal="center")
        if isinstance(periode, date):
            feuille.cell(ligne_periode, 1).number_format = "yyyy-mm-dd"
        total = index == len(blocs) - 1
        fond = "EDF1F6" if total else "78A3CF" if index % 2 == 0 else "3768A0"
        texte = "152331" if total or index % 2 == 0 else "FFFFFF"
        for categorie in categories:
            _ajouter_ligne(feuille, [categorie.libelle, categorie.enregistrees, categorie.annulees, categorie.net])
            for cellule in feuille[feuille.max_row]:
                cellule.fill = PatternFill("solid", fgColor=fond)
                cellule.font = Font(bold=total, color=texte)
                cellule.alignment = Alignment(horizontal="left" if cellule.column == 1 else "right")
    _mettre_en_forme(feuille, filtre=False)


def exporter_excel(rapport):
    classeur = Workbook()
    parametres = classeur.active
    parametres.title = "Paramètres"
    _ajouter_ligne(parametres, ["Paramètre", "Valeur"])
    for nom, valeur in (
        ("Date de début", rapport.date_debut), ("Date de fin", rapport.date_fin),
        ("Niveau de détail", dict(PRECISIONS)[rapport.precision]),
        ("Types de transactions", ", ".join(dict(TYPES_RAPPORT)[type_cafe] for type_cafe in rapport.types)),
        ("Fuseau horaire", str(FUSEAU_RAPPORT)),
        ("Généré le", rapport.genere_le.replace(tzinfo=None)),
        ("Annulations", "Déduites à leur date d’annulation"),
    ):
        _ajouter_ligne(parametres, [nom, valeur])
    _mettre_en_forme(parametres, filtre=False)

    synthese = classeur.create_sheet("Synthèse")
    entete_periode = "Date" if rapport.precision == "jour" else "Période"
    _ajouter_ligne(synthese, [entete_periode, *(categorie.libelle for categorie in rapport.totaux), "TOTAL"])
    for groupe in rapport.groupes:
        periode = groupe.debut if groupe.debut == groupe.fin else groupe.libelle_periode
        _ajouter_ligne(synthese, [periode, *(categorie.net for categorie in groupe.categories), groupe.total])
    _ajouter_ligne(synthese, ["TOTAL", *(categorie.net for categorie in rapport.totaux), rapport.total])
    _mettre_en_forme(synthese)
    synthese.freeze_panes = "B2"
    synthese.auto_filter.ref = f"A1:{get_column_letter(synthese.max_column)}{synthese.max_row - 1}"
    for cellule in synthese[synthese.max_row]:
        cellule.font = Font(bold=True)
        cellule.fill = PatternFill("solid", fgColor="EDF1F6")

    _detail_totaux_excel(classeur, rapport)

    if rapport.precision == "transaction":
        transactions = classeur.create_sheet("Transactions")
        _ajouter_ligne(transactions, ["Date", "Transaction", "Type", "Événement", "Client", "Employé",
                                      "Enregistrées", "Annulées", "Net"])
        for evenement in rapport.evenements:
            _ajouter_ligne(transactions, [evenement.date.replace(tzinfo=None), evenement.transaction_id,
                                         evenement.type_transaction, evenement.nature, evenement.client,
                                         evenement.employe, evenement.enregistrees, evenement.annulees, evenement.net])
        _mettre_en_forme(transactions)
    fichier = BytesIO()
    classeur.save(fichier)
    fichier.seek(0)
    return fichier
