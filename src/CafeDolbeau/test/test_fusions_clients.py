from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict
from threading import Barrier
from unittest.mock import patch

from django.contrib import admin
from django.contrib.admin.models import CHANGE, DELETION, LogEntry
from django.core import signing
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import close_old_connections, connection, transaction
from django.db.migrations.executor import MigrationExecutor
from django.db.models.signals import pre_delete
from django.test import Client as Navigateur, TestCase, TransactionTestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from ..forms.forms_annulation import SelectionAnciennesForm
from ..models import AnnulationTransaction, Client, Employe, FusionClient, TransactionCafe
from ..permissions import peut_annuler
from ..rapports import exporter_excel, generer_rapport
from ..services import COMPTEURS, ajouter_cafes, annuler_fusion, annuler_operation, fusionner_clients, supprimer_clients
from .test_annulations import creer_employe


def compteurs(personne):
    personne.refresh_from_db()
    return {champ: getattr(personne, champ) for champ, _, _ in COMPTEURS}


class FusionClientTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.gerant = creer_employe("gerant-fusion", "Gérant")
        cls.barista = creer_employe("barista-fusion")
        cls.source = Client.objects.create(
            nom_complet="Camille", telephone="418 555-1000", nombre_cafes_achetes=7,
            nombre_cafes_prepayes=3, nombre_cafes_gratuits=2,
        )
        cls.destination = Client.objects.create(
            nom_complet="Alex", courriel="alex@example.com", nombre_cafes_achetes=9,
            nombre_cafes_prepayes=5, nombre_cafes_gratuits=4,
        )
        cls.autre = Client.objects.create(nom_complet="Lou", telephone="418 555-3000")

    def setUp(self):
        self.client.force_login(self.gerant)
        self.liste = reverse("admin:CafeDolbeau_client_changelist")

    def fusionner(self, source=None, destination=None, employe=None):
        return fusionner_clients(
            (source or self.source).pk, (destination or self.destination).pk,
            employe=employe or self.gerant,
        )

    def acheter(self, personne=None, quantite=1, type_cafe="achat"):
        return ajouter_cafes(personne or self.destination, quantite, type_cafe, employe=self.barista)

    def ancienne(self, personne=None, origine=None):
        return TransactionCafe.objects.create(
            client=personne or self.source, quantite=1, type_transaction="achat", client_origine_id=origine,
        )

    def action(self, **donnees):
        return self.client.post(self.liste, {
            "action": "fusionner_selection", "_selected_action": [self.source.pk, self.destination.pk], **donnees,
        })

    def apercu(self):
        reponse = self.action(etape="previsualiser", destination=self.destination.pk)
        self.assertEqual(reponse.status_code, 200)
        return reponse

    def test_transfert_addition_sans_recompense_et_coordonnees_conservees(self):
        self.acheter(self.source)
        self.acheter()
        avant = list(TransactionCafe.objects.order_by("pk").values())
        initiaux = compteurs(self.source), compteurs(self.destination)
        fusion = self.fusionner()
        self.assertEqual(compteurs(self.source), initiaux[0])
        self.assertFalse(self.source.est_actif)
        self.assertEqual(compteurs(self.destination), {champ: initiaux[0][champ] + initiaux[1][champ] for champ, _, _ in COMPTEURS})
        self.assertEqual((self.destination.nom_complet, self.destination.courriel, self.destination.telephone), ("Alex", "alex@example.com", ""))
        for ligne, initiale in zip(TransactionCafe.objects.order_by("pk").values(), avant):
            attendue = initiale if initiale["client_id"] == self.destination.pk else {
                **initiale, "client_id": self.destination.pk, "client_origine_id": self.source.pk,
            }
            self.assertEqual(ligne, attendue)
        self.assertEqual(fusion.statut, FusionClient.Statut.EN_ATTENTE)
        self.assertEqual((fusion.cree_par_id, str(fusion.client_source), str(fusion.client_destination)), (self.gerant.pk, "Camille", "Alex"))
        self.assertEqual(LogEntry.objects.filter(action_flag=CHANGE).count(), 2)
        self.assertFalse(Client.objects.filter(pk=-1).exists())

    def test_retour_arriere_exact_provenances_et_annulations_incluses(self):
        annulee = self.acheter(self.source)[0]
        annuler_operation(annulee.pk, self.gerant)
        self.ancienne(origine=12345)
        self.ancienne()
        self.ancienne(self.destination)
        clients = list(Client.objects.order_by("pk").values())
        lignes = list(TransactionCafe.objects.order_by("pk").values())
        annulations = list(AnnulationTransaction.objects.values())
        fusion = self.fusionner()
        self.assertIn({"id": lignes[1]["id"], "client_origine_id": 12345}, fusion.transactions_deplacees)
        restauree = annuler_fusion(fusion.pk, employe=self.gerant)
        self.assertEqual(list(Client.objects.order_by("pk").values()), clients)
        self.assertEqual(list(TransactionCafe.objects.order_by("pk").values()), lignes)
        self.assertEqual(list(AnnulationTransaction.objects.values()), annulations)
        self.assertEqual(restauree.statut, FusionClient.Statut.ANNULEE)
        self.assertEqual(restauree.terminee_par_id, self.gerant.pk)
        self.assertIsNone(restauree.operation_finalisation_id)
        self.assertEqual(LogEntry.objects.count(), 4)
        annuler_fusion(fusion.pk, employe=self.gerant)
        self.assertEqual(LogEntry.objects.count(), 4)

    def test_sans_transactions_retour_arriere_et_finalisation(self):
        fusion = self.fusionner()
        self.assertEqual(fusion.transactions_deplacees, [])
        annuler_fusion(fusion.pk, employe=self.gerant)
        fusion = self.fusionner()
        lignes = self.acheter()
        fusion.refresh_from_db()
        self.assertEqual(fusion.statut, FusionClient.Statut.FINALISEE)
        self.assertEqual(fusion.operation_finalisation_id, lignes[0].operation_id)

    def test_trois_types_operation_finalisent_seulement_apres_enregistrement(self):
        for type_cafe in ("achat", "prepaye", "utilise"):
            with self.subTest(type_cafe=type_cafe), transaction.atomic():
                source = Client.objects.create(nom_complet=type_cafe, telephone=f"source-{type_cafe}")
                destination = Client.objects.create(nom_complet=type_cafe, telephone=f"dest-{type_cafe}", nombre_cafes_prepayes=5)
                fusion = self.fusionner(source, destination)
                lignes = self.acheter(destination, type_cafe=type_cafe)
                self.assertTrue(lignes)
                fusion.refresh_from_db()
                self.assertEqual(fusion.statut, FusionClient.Statut.FINALISEE)
                self.assertEqual(fusion.terminee_par_id, self.barista.pk)
                self.assertEqual(fusion.operation_finalisation_id, lignes[0].operation_id)
                self.assertFalse(Client.objects.filter(pk=source.pk).exists())
                self.assertFalse(Client.objects.filter(pk=-1).exists())
                self.assertEqual(LogEntry.objects.filter(object_id=str(source.pk), action_flag=DELETION).count(), 1)

    def test_autres_clients_sans_effet_et_annulation_declencheur_ne_rouvre_pas(self):
        fusion = self.fusionner()
        self.acheter(self.autre)
        fusion.refresh_from_db()
        self.assertEqual(fusion.statut, FusionClient.Statut.EN_ATTENTE)
        lignes = self.acheter()
        annuler_operation(lignes[0].pk, self.gerant)
        fusion.refresh_from_db()
        self.assertEqual(fusion.statut, FusionClient.Statut.FINALISEE)
        self.assertFalse(Client.objects.filter(pk=self.source.pk).exists())
        with self.assertRaises(ValidationError):
            annuler_fusion(fusion.pk, employe=self.gerant)

    def test_historique_survit_suppression_des_participants(self):
        self.ancienne()
        fusion = self.fusionner()
        self.acheter()
        supprimer_clients([self.destination.pk], employe=self.gerant)
        fusion.refresh_from_db()
        self.assertIsNone(fusion.client_source_id)
        self.assertIsNone(fusion.client_destination_id)
        self.assertEqual(str(fusion), "Client supprimé → Client supprimé")
        self.assertContains(self.client.get(reverse("admin:CafeDolbeau_fusionclient_change", args=[fusion.pk])), "Client supprimé")

    def test_noms_affiches_et_recherche_suivent_les_clients(self):
        fusion = self.fusionner()
        Client.objects.filter(pk=self.source.pk).update(nom_complet="Camille corrigé")
        Client.objects.filter(pk=self.destination.pk).update(nom_complet="Alex corrigé")
        fusion.refresh_from_db()
        self.assertEqual(str(fusion), "Camille corrigé → Alex corrigé")
        liste = reverse("admin:CafeDolbeau_fusionclient_changelist")
        for recherche in ("Camille corrigé", "Alex corrigé", str(self.source.pk), str(self.destination.pk)):
            reponse = self.client.get(liste, {"q": recherche})
            self.assertEqual([ligne.pk for ligne in reponse.context["cl"].result_list], [fusion.pk])
        for url in (
            reverse("admin:CafeDolbeau_fusionclient_change", args=[fusion.pk]),
            reverse("admin:CafeDolbeau_fusionclient_annuler", args=[fusion.pk]),
        ):
            reponse = self.client.get(url)
            self.assertContains(reponse, "Camille corrigé")
            self.assertContains(reponse, "Alex corrigé")
        self.assertContains(self.client.get(reverse("ajouter_cafes", args=[self.destination.pk])), "Camille corrigé")

    def test_suppression_participant_fusion_annulee_conserve_historique(self):
        for personne, champ, autre_champ, autre_id in (
            (self.source, "client_source_id", "client_destination_id", self.destination.pk),
            (self.destination, "client_destination_id", "client_source_id", self.source.pk),
        ):
            with self.subTest(client=personne.pk), transaction.atomic():
                fusion = self.fusionner()
                annuler_fusion(fusion.pk, employe=self.gerant)
                self.assertEqual(supprimer_clients([personne.pk], employe=self.gerant), 1)
                fusion.refresh_from_db()
                self.assertIsNone(getattr(fusion, champ))
                self.assertEqual(getattr(fusion, autre_champ), autre_id)
                self.assertEqual(fusion.statut, FusionClient.Statut.ANNULEE)
                self.assertTrue(Client.objects.filter(pk=autre_id).exists())
                transaction.set_rollback(True)

    def test_suppression_en_lot_conserve_fusion_sans_participants(self):
        fusion = self.fusionner()
        annuler_fusion(fusion.pk, employe=self.gerant)
        self.assertEqual(supprimer_clients([self.source.pk, self.destination.pk], employe=self.gerant), 2)
        fusion.refresh_from_db()
        self.assertIsNone(fusion.client_source_id)
        self.assertIsNone(fusion.client_destination_id)
        self.assertEqual(fusion.statut, FusionClient.Statut.ANNULEE)
        self.assertEqual(str(fusion), "Client supprimé → Client supprimé")

    def test_permissions_fraiches_fusion_et_retour_arriere(self):
        with self.assertRaises(PermissionDenied):
            self.fusionner(employe=self.barista)
        fusion = self.fusionner()
        with self.assertRaises(PermissionDenied):
            annuler_fusion(fusion.pk, employe=self.barista)
        self.assertTrue(self.gerant.has_perm("CafeDolbeau.acces_gerant"))
        self.gerant.groups.clear()
        with self.assertRaises(PermissionDenied):
            annuler_fusion(fusion.pk, employe=self.gerant)
        with self.assertRaises(PermissionDenied):
            fusionner_clients(self.autre.pk, self.destination.pk, employe=self.gerant)

    def test_superutilisateur_sans_groupe(self):
        superuser = Employe.objects.create_superuser("super-fusion", "super@example.com", "test")
        fusion = self.fusionner(employe=superuser)
        self.client.force_login(superuser)
        self.assertEqual(self.client.get(reverse("admin:CafeDolbeau_fusionclient_changelist")).status_code, 200)
        self.assertEqual(annuler_fusion(fusion.pk, employe=superuser).statut, FusionClient.Statut.ANNULEE)

    def test_clients_invalides(self):
        for source, destination in ((self.source.pk, self.source.pk), (-1, self.destination.pk), (self.source.pk, -1), (999999, self.destination.pk), (None, self.destination.pk)):
            with self.subTest(source=source, destination=destination), self.assertRaises(ValidationError):
                fusionner_clients(source, destination, employe=self.gerant)
        Client.objects.filter(pk=self.source.pk).update(est_actif=False)
        with self.assertRaises(ValidationError):
            self.fusionner()
        self.assertEqual(FusionClient.objects.count(), 0)

    def test_limites_des_trois_compteurs(self):
        for champ, _, _ in COMPTEURS:
            with self.subTest(champ=champ), transaction.atomic():
                Client.objects.filter(pk=self.source.pk).update(**{champ: 2147483647})
                with self.assertRaises(ValidationError):
                    self.fusionner()
                transaction.set_rollback(True)
        self.assertFalse(FusionClient.objects.exists())

    def test_une_seule_fusion_en_attente_par_participant(self):
        self.fusionner()
        for source, destination in ((self.source, self.autre), (self.destination, self.autre), (self.autre, self.destination), (self.autre, self.source)):
            with self.subTest(source=source.pk, destination=destination.pk), self.assertRaises(ValidationError):
                self.fusionner(source, destination)
        self.assertEqual(FusionClient.objects.count(), 1)

    def test_protections_suppression_achat_source_et_annulation_transactions(self):
        ligne = self.acheter(self.source)[0]
        self.fusionner()
        ligne.refresh_from_db()
        self.assertFalse(peut_annuler(ligne, self.gerant))
        with self.assertRaises(ValidationError):
            annuler_operation(ligne.pk, self.gerant)
        with self.assertRaises(ValidationError):
            self.acheter(self.source)
        for personne in (self.source, self.destination):
            with self.subTest(personne=personne.pk), self.assertRaises(ValidationError):
                supprimer_clients([personne.pk], employe=self.gerant)
        self.assertEqual(self.client.get(reverse("admin:CafeDolbeau_client_delete", args=[self.destination.pk])).status_code, 403)

    def test_confirmation_annulation_anterieure_et_formulaire_ancien_bloques(self):
        ligne = self.acheter(self.source)[0]
        ancienne = self.ancienne(self.destination)
        url = reverse("annuler_transaction_admin", args=[ligne.pk])
        token = self.client.get(url).context["confirmation"]
        self.fusionner()
        self.assertEqual(self.client.post(url, {"action": "confirmer", "confirmation": token}).status_code, 400)
        reponse = self.client.get(reverse("annuler_transaction_admin", args=[ancienne.pk]))
        self.assertContains(reponse, "Annulez ou finalisez la fusion", status_code=400)
        self.assertNotIn("selection", reponse.context)
        self.assertFalse(AnnulationTransaction.objects.exists())

    def test_operations_transferees_restent_groupees_apres_finalisation(self):
        Client.objects.filter(pk=self.source.pk).update(nombre_cafes_achetes=9)
        deplacees = self.acheter(self.source)
        self.assertEqual(len(deplacees), 2)
        self.fusionner()
        nouvelles = self.acheter()
        self.assertEqual(annuler_operation(deplacees[0].pk, self.gerant), 2)
        self.assertEqual(set(AnnulationTransaction.objects.values_list("transaction_id", flat=True)), {ligne.pk for ligne in deplacees})
        self.assertFalse(AnnulationTransaction.objects.filter(transaction_id__in=[ligne.pk for ligne in nouvelles]).exists())

    def test_protection_orm_des_clients_sans_transactions_et_historique(self):
        fusion = self.fusionner()
        for personne in (self.source, self.destination):
            with self.assertRaises(ValidationError), transaction.atomic():
                Client.objects.filter(pk=personne.pk).delete()
        with self.assertRaises(ValidationError), transaction.atomic():
            FusionClient.objects.filter(pk=fusion.pk).delete()

    def test_fusions_successives_et_provenances_preservees_apres_suppression(self):
        ancienne = self.ancienne(origine=12345)
        ligne = self.ancienne()
        fusion = self.fusionner()
        self.acheter()
        seconde = self.fusionner(self.destination, self.autre)
        snapshots = list(TransactionCafe.objects.order_by("pk").values())
        annuler_fusion(seconde.pk, employe=self.gerant)
        for original in snapshots:
            ligne_actuelle = TransactionCafe.objects.get(pk=original["id"])
            self.assertEqual(ligne_actuelle.client_id, self.destination.pk)
        seconde = self.fusionner(self.destination, self.autre)
        self.acheter(self.autre)
        supprimer_clients([self.autre.pk], employe=self.gerant)
        for tx, origine in ((ancienne, 12345), (ligne, self.source.pk)):
            tx.refresh_from_db()
            self.assertEqual((tx.client_id, tx.client_origine_id), (-1, origine))
        self.assertEqual(FusionClient.objects.count(), 3)
        fusion.refresh_from_db()
        self.assertEqual(fusion.statut, FusionClient.Statut.FINALISEE)

    def test_anciennes_transactions_selection_limitee_par_provenance(self):
        source = self.ancienne()
        destination = self.ancienne(self.destination)
        self.fusionner()
        self.acheter()
        source.refresh_from_db()
        form = SelectionAnciennesForm(client=self.destination, source=source)
        self.assertEqual(set(form.fields["lignes"].queryset.values_list("pk", flat=True)), {source.pk})
        avec_melange = SelectionAnciennesForm({"lignes": [source.pk, destination.pk]}, client=self.destination, source=source)
        self.assertFalse(avec_melange.is_valid())
        with self.assertRaises(ValidationError):
            annuler_operation(source.pk, self.gerant, {source.pk: (1, 0, 0), destination.pk: (1, 0, 0)})
        self.assertEqual(annuler_operation(source.pk, self.gerant, {source.pk: (1, 0, 0)}), 1)
        self.assertFalse(TransactionCafe.objects.get(pk=destination.pk).annulee)

    def test_provenance_connue_et_inconnue_du_meme_client_selectionnables(self):
        source = self.ancienne()
        source_connue = self.ancienne(origine=self.source.pk)
        self.fusionner()
        self.acheter()
        source.refresh_from_db()
        self.assertEqual(annuler_operation(source.pk, self.gerant, {source.pk: (1, 0, 0), source_connue.pk: (1, 0, 0)}), 2)

    def test_erreur_journal_fusion_annule_tous_les_changements(self):
        self.ancienne()
        clients = list(Client.objects.values())
        lignes = list(TransactionCafe.objects.values())
        with patch("CafeDolbeau.services.LogEntry.objects.log_actions", side_effect=RuntimeError("journal")), self.assertRaises(RuntimeError):
            self.fusionner()
        self.assertEqual(list(Client.objects.values()), clients)
        self.assertEqual(list(TransactionCafe.objects.values()), lignes)
        self.assertFalse(FusionClient.objects.exists())

    def test_erreur_journal_retour_arriere_preserve_fusion_en_attente(self):
        self.ancienne()
        fusion = self.fusionner()
        clients, lignes = list(Client.objects.values()), list(TransactionCafe.objects.values())
        with patch("CafeDolbeau.services.LogEntry.objects.log_actions", side_effect=RuntimeError("journal")), self.assertRaises(RuntimeError):
            annuler_fusion(fusion.pk, employe=self.gerant)
        self.assertEqual(list(Client.objects.values()), clients)
        self.assertEqual(list(TransactionCafe.objects.values()), lignes)
        fusion.refresh_from_db()
        self.assertEqual(fusion.statut, FusionClient.Statut.EN_ATTENTE)

    def test_erreur_suppression_finale_annule_achat_compteurs_et_finalisation(self):
        fusion = self.fusionner()
        clients, lignes = list(Client.objects.values()), list(TransactionCafe.objects.values())

        def echouer(sender, instance, **kwargs):
            if instance.pk == self.source.pk:
                raise RuntimeError("suppression")

        pre_delete.connect(echouer, sender=Client, weak=False)
        try:
            with self.assertRaises(RuntimeError):
                self.acheter()
        finally:
            pre_delete.disconnect(echouer, sender=Client)
        self.assertEqual(list(Client.objects.values()), clients)
        self.assertEqual(list(TransactionCafe.objects.values()), lignes)
        fusion.refresh_from_db()
        self.assertEqual(fusion.statut, FusionClient.Statut.EN_ATTENTE)
        self.assertIsNone(fusion.operation_finalisation_id)
        self.assertFalse(LogEntry.objects.filter(action_flag=DELETION).exists())

    def test_operation_invalide_et_depassement_ne_finalisent_pas(self):
        Client.objects.filter(pk=self.source.pk).update(nombre_cafes_achetes=0)
        Client.objects.filter(pk=self.destination.pk).update(nombre_cafes_achetes=2147483647)
        fusion = self.fusionner()
        for quantite in (0, 1):
            with self.assertRaises(ValidationError):
                self.acheter(quantite=quantite)
        fusion.refresh_from_db()
        self.assertEqual(fusion.statut, FusionClient.Statut.EN_ATTENTE)
        self.assertTrue(Client.objects.filter(pk=self.source.pk).exists())
        self.assertEqual(TransactionCafe.objects.count(), 0)

    def test_rapport_et_export_conservent_evenements_totaux_et_annulations(self):
        ligne = self.acheter(self.source)[0]
        annuler_operation(ligne.pk, self.gerant)
        self.acheter(self.source)
        date = timezone.localdate()
        parametres = dict(date_debut=date, date_fin=date, types=[valeur for valeur, _ in TransactionCafe.Type.choices], precision="transaction")
        avant = generer_rapport(**parametres)
        fusion = self.fusionner()
        apres = generer_rapport(**parametres)
        self.assertEqual([asdict(total) for total in avant.totaux], [asdict(total) for total in apres.totaux])
        self.assertEqual([asdict(tx) for tx in apres.evenements], [{**asdict(tx), "client": "Alex"} for tx in avant.evenements])
        excel = load_workbook(exporter_excel(apres))
        self.assertEqual({ligne[4] for ligne in excel["Transactions"].iter_rows(min_row=2, values_only=True)}, {"Alex"})
        annuler_fusion(fusion.pk, employe=self.gerant)
        self.assertEqual([asdict(tx) for tx in generer_rapport(**parametres).evenements], [asdict(tx) for tx in avant.evenements])

    def test_action_exige_exactement_deux_et_choix_explicite(self):
        for selection in ([self.source.pk], [self.source.pk, self.destination.pk, self.autre.pk]):
            self.assertEqual(self.action(_selected_action=selection).status_code, 302)
        reponse = self.action()
        self.assertContains(reponse, "Client à conserver")
        self.assertNotIn("confirmation", reponse.context)
        self.assertFalse(FusionClient.objects.exists())
        reponse = self.action(etape="previsualiser")
        self.assertTrue(reponse.context["form"].errors)
        self.assertFalse(FusionClient.objects.exists())

    def test_confirmation_ne_modifie_rien_et_post_fusionne(self):
        clients = list(Client.objects.values())
        reponse = self.apercu()
        self.assertContains(reponse, "sans récompense supplémentaire")
        self.assertEqual(list(Client.objects.values()), clients)
        self.assertFalse(FusionClient.objects.exists())
        self.assertFalse(LogEntry.objects.exists())
        self.assertEqual(self.action(etape="confirmer", destination=self.destination.pk, confirmation=reponse.context["confirmation"]).status_code, 302)
        self.assertEqual(FusionClient.objects.count(), 1)

    def test_confirmation_perimee_presente_nouveaux_soldes_sans_fusion(self):
        reponse = self.apercu()
        self.acheter(self.source)
        reponse = self.action(etape="confirmer", destination=self.destination.pk, confirmation=reponse.context["confirmation"])
        self.assertEqual(reponse.status_code, 409)
        self.assertContains(reponse, "Les soldes ont changé", status_code=409)
        self.assertFalse(FusionClient.objects.exists())
        self.assertEqual(self.action(etape="confirmer", destination=self.destination.pk, confirmation=reponse.context["confirmation"]).status_code, 302)

    def test_confirmation_invalide_autre_destination_ou_autre_employe(self):
        token = self.apercu().context["confirmation"]
        for donnees in ({"confirmation": "invalide"}, {"confirmation": token, "destination": self.source.pk}):
            donnees = {"destination": self.destination.pk, **donnees}
            self.assertEqual(self.action(etape="confirmer", **donnees).status_code, 400)
        autre = creer_employe("autre-gerant-fusion", "Gérant")
        self.client.force_login(autre)
        self.assertEqual(self.action(etape="confirmer", destination=self.destination.pk, confirmation=token).status_code, 400)
        self.assertFalse(FusionClient.objects.exists())

    def test_confirmation_si_client_supprime_ou_deja_fusionne(self):
        token = self.apercu().context["confirmation"]
        self.fusionner()
        reponse = self.action(etape="confirmer", destination=self.destination.pk, confirmation=token)
        self.assertEqual(reponse.status_code, 400)
        self.assertEqual(FusionClient.objects.count(), 1)
        self.assertNotIn("confirmation", reponse.context)

    def test_action_et_historique_reserves_aux_gerants(self):
        # Même un barista auquel l'accès technique à l'admin serait accordé.
        Employe.objects.filter(pk=self.barista.pk).update(is_staff=True)
        self.client.force_login(self.barista)
        self.assertEqual(self.client.get(reverse("admin:CafeDolbeau_fusionclient_changelist"), follow=True).status_code, 403)
        self.action(etape="previsualiser", destination=self.destination.pk)
        self.assertFalse(FusionClient.objects.exists())
        self.assertNotContains(self.client.get(self.liste, follow=True), "Fusionner deux clients", status_code=403)

    def test_source_masquee_redirection_get_et_post_refuse(self):
        self.fusionner()
        url = reverse("ajouter_cafes", args=[self.source.pk])
        self.assertNotContains(self.client.get(reverse("accueil")), "Camille")
        self.assertRedirects(self.client.get(url), reverse("ajouter_cafes", args=[self.destination.pk]))
        self.assertEqual(self.client.post(url, {"quantite": 1, "type_transaction": "achat"}).status_code, 403)
        self.assertFalse(TransactionCafe.objects.exists())

    def test_avertissement_fiche_disparait_apres_operation_ajax(self):
        self.fusionner()
        url = reverse("ajouter_cafes", args=[self.destination.pk])
        self.assertContains(self.client.get(url), "la prochaine opération enregistrée rendra la fusion définitive")
        reponse = self.client.post(url, {"quantite": 1, "type_transaction": "achat"}, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(reponse.status_code, 200)
        self.assertFalse(reponse.json()["fusion_en_attente"])
        self.assertNotContains(self.client.get(url), "avertissement-fusion")

    def test_fiches_admin_source_lecture_seule_et_compteurs_destination_figes(self):
        self.fusionner()
        source_url = reverse("admin:CafeDolbeau_client_change", args=[self.source.pk])
        destination_url = reverse("admin:CafeDolbeau_client_change", args=[self.destination.pk])
        source = self.client.get(source_url)
        self.assertEqual(source.status_code, 200)
        self.assertNotContains(source, 'name="nom_complet"')
        self.assertEqual(self.client.post(source_url, {"nom_complet": "modifié"}).status_code, 403)
        destination = self.client.get(destination_url)
        for champ, _, _ in COMPTEURS:
            self.assertIn(champ, destination.context["adminform"].readonly_fields)
        initiaux = compteurs(self.destination)
        reponse = self.client.post(destination_url, {
            "nom_complet": "Alex corrigé", "telephone": "", "courriel": "alex@example.com",
            "nombre_cafes_achetes": 999, "nombre_cafes_prepayes": 999, "nombre_cafes_gratuits": 999,
        })
        self.assertEqual(reponse.status_code, 302)
        self.assertEqual(compteurs(self.destination), initiaux)
        self.assertEqual(self.destination.nom_complet, "Alex corrigé")

    def test_modification_coordonnees_necrase_pas_operation_intervenue_depuis_lecture(self):
        fusion = self.fusionner()
        modeladmin = admin.site._registry[Client]
        enregistrer = modeladmin.save_model

        def acheter_avant_enregistrement(request, obj, form, change):
            self.acheter()
            enregistrer(request, obj, form, change)

        with patch.object(modeladmin, "save_model", side_effect=acheter_avant_enregistrement):
            reponse = self.client.post(reverse("admin:CafeDolbeau_client_change", args=[self.destination.pk]), {
                "nom_complet": "Alex corrigé", "telephone": "", "courriel": "alex@example.com",
            })
        self.assertEqual(reponse.status_code, 302)
        fusion.refresh_from_db()
        self.assertEqual(fusion.statut, FusionClient.Statut.FINALISEE)
        self.assertEqual(compteurs(self.destination)["nombre_cafes_achetes"], 17)
        self.assertEqual(self.destination.nom_complet, "Alex corrigé")

    def test_historique_interdit_modification_et_suppression(self):
        fusion = self.fusionner()
        self.assertEqual(self.client.get(reverse("admin:CafeDolbeau_fusionclient_add")).status_code, 403)
        url = reverse("admin:CafeDolbeau_fusionclient_change", args=[fusion.pk])
        self.assertEqual(self.client.post(url, {"statut": "finalisee"}).status_code, 403)
        self.assertEqual(self.client.get(reverse("admin:CafeDolbeau_fusionclient_delete", args=[fusion.pk])).status_code, 403)
        self.assertContains(self.client.get(url), "Annuler cette fusion")

    def test_confirmation_retour_arriere_sans_ecriture_et_post(self):
        fusion = self.fusionner()
        url = reverse("admin:CafeDolbeau_fusionclient_annuler", args=[fusion.pk])
        clients = list(Client.objects.values())
        reponse = self.client.get(url)
        self.assertContains(reponse, "Confirmer le retour arrière")
        self.assertEqual(list(Client.objects.values()), clients)
        self.assertEqual(self.client.post(url, {"confirmation": reponse.context["confirmation"]}).status_code, 302)
        fusion.refresh_from_db()
        self.assertEqual(fusion.statut, FusionClient.Statut.ANNULEE)

    def test_confirmation_retour_arriere_devenue_definitive(self):
        fusion = self.fusionner()
        url = reverse("admin:CafeDolbeau_fusionclient_annuler", args=[fusion.pk])
        token = self.client.get(url).context["confirmation"]
        self.acheter()
        reponse = self.client.post(url, {"confirmation": token})
        self.assertEqual(reponse.status_code, 409)
        self.assertNotIn("confirmation", reponse.context)
        self.assertContains(reponse, "définitive", status_code=409)

    def test_retour_arriere_invalide_barista_et_csrf(self):
        fusion = self.fusionner()
        url = reverse("admin:CafeDolbeau_fusionclient_annuler", args=[fusion.pk])
        self.assertEqual(self.client.post(url, {"confirmation": "invalide"}).status_code, 400)
        navigateur = Navigateur(enforce_csrf_checks=True)
        navigateur.force_login(self.gerant)
        self.assertEqual(navigateur.post(url, {}).status_code, 403)
        self.assertEqual(navigateur.post(self.liste, {}).status_code, 403)
        Employe.objects.filter(pk=self.barista.pk).update(is_staff=True)
        self.client.force_login(self.barista)
        self.assertEqual(self.client.get(url, follow=True).status_code, 403)


class FusionConcurrenceTests(TransactionTestCase):
    def setUp(self):
        self.gerant = creer_employe("gerant-fusion-concurrence", "Gérant")
        self.barista = creer_employe("barista-fusion-concurrence")
        self.source = Client.objects.create(nom_complet="Source", telephone="111", nombre_cafes_achetes=2)
        self.destination = Client.objects.create(nom_complet="Destination", telephone="222", nombre_cafes_achetes=3)
        self.autre = Client.objects.create(nom_complet="Autre", telephone="333")

    def simultanement(self, fonctions):
        barriere = Barrier(len(fonctions))

        def executer(fonction):
            close_old_connections()
            try:
                barriere.wait(timeout=10)
                try:
                    return fonction()
                except (ValidationError, Client.DoesNotExist) as erreur:
                    return erreur
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=len(fonctions)) as pool:
            return list(pool.map(executer, fonctions))

    def fusionner(self):
        return fusionner_clients(self.source.pk, self.destination.pk, employe=self.gerant)

    def acheter(self, personne=None):
        return ajouter_cafes(personne or self.destination, 1, employe=self.barista)

    def test_achat_et_annulation_fusion_concurrents(self):
        fusion = self.fusionner()
        lignes, resultat = self.simultanement([self.acheter, lambda: annuler_fusion(fusion.pk, employe=self.gerant)])
        fusion.refresh_from_db()
        self.assertEqual(TransactionCafe.objects.count(), len(lignes))
        if fusion.statut == FusionClient.Statut.FINALISEE:
            self.assertIsInstance(resultat, ValidationError)
            self.assertFalse(Client.objects.filter(pk=self.source.pk).exists())
            self.assertEqual(compteurs(self.destination)["nombre_cafes_achetes"], 6)
        else:
            self.assertEqual(fusion.statut, FusionClient.Statut.ANNULEE)
            self.assertEqual(compteurs(self.source)["nombre_cafes_achetes"], 2)
            self.assertEqual(compteurs(self.destination)["nombre_cafes_achetes"], 4)

    def test_deux_achats_finalisent_une_seule_fois(self):
        fusion = self.fusionner()
        self.simultanement([self.acheter, self.acheter])
        fusion.refresh_from_db()
        self.assertEqual(fusion.statut, FusionClient.Statut.FINALISEE)
        self.assertEqual(TransactionCafe.objects.count(), 2)
        self.assertEqual(compteurs(self.destination)["nombre_cafes_achetes"], 7)
        self.assertEqual(LogEntry.objects.filter(action_flag=DELETION).count(), 1)

    def test_deux_fusions_simultanees_un_seul_historique(self):
        resultats = self.simultanement([self.fusionner, self.fusionner])
        self.assertEqual(sum(isinstance(resultat, FusionClient) for resultat in resultats), 1)
        self.assertEqual(FusionClient.objects.count(), 1)
        self.assertEqual(compteurs(self.destination)["nombre_cafes_achetes"], 5)

    def test_participant_commun_roles_croises(self):
        resultats = self.simultanement([
            self.fusionner, lambda: fusionner_clients(self.destination.pk, self.autre.pk, employe=self.gerant),
        ])
        self.assertEqual(sum(isinstance(resultat, FusionClient) for resultat in resultats), 1)
        self.assertEqual(FusionClient.objects.count(), 1)

    def test_achat_source_et_fusion_serialises(self):
        fusion, achat = self.simultanement([self.fusionner, lambda: self.acheter(self.source)])
        self.assertEqual(fusion.statut, FusionClient.Statut.EN_ATTENTE)
        quantite = 0 if isinstance(achat, ValidationError) else 1
        self.assertEqual(compteurs(self.destination)["nombre_cafes_achetes"], 5 + quantite)
        self.assertEqual(TransactionCafe.objects.count(), quantite)
        self.assertFalse(TransactionCafe.objects.exclude(client=self.destination).exists())

    def test_annulation_transaction_et_fusion_serialisees(self):
        ligne = self.acheter(self.source)[0]
        fusion, annulation = self.simultanement([self.fusionner, lambda: annuler_operation(ligne.pk, self.gerant)])
        ligne.refresh_from_db()
        self.assertEqual(ligne.client_id, self.destination.pk)
        annulee = not isinstance(annulation, ValidationError)
        self.assertEqual(ligne.annulee, annulee)
        self.assertEqual(compteurs(self.destination)["nombre_cafes_achetes"], 5 if annulee else 6)
        self.assertEqual(fusion.statut, FusionClient.Statut.EN_ATTENTE)

    def test_suppression_et_fusion_serialisees(self):
        fusion, suppression = self.simultanement([
            self.fusionner, lambda: supprimer_clients([self.source.pk], employe=self.gerant),
        ])
        if isinstance(fusion, FusionClient):
            self.assertIsInstance(suppression, ValidationError)
            self.assertTrue(Client.objects.filter(pk=self.source.pk).exists())
        else:
            self.assertEqual(suppression, 1)
            self.assertEqual(FusionClient.objects.count(), 0)

    def test_deux_retours_arriere_ne_dupliquent_pas_la_restitution(self):
        fusion = self.fusionner()
        self.simultanement([lambda: annuler_fusion(fusion.pk, employe=self.gerant)] * 2)
        self.assertEqual(compteurs(self.source)["nombre_cafes_achetes"], 2)
        self.assertEqual(compteurs(self.destination)["nombre_cafes_achetes"], 3)
        self.assertEqual(LogEntry.objects.count(), 4)


class FusionMigrationTests(TransactionTestCase):
    def test_migration_retire_noms_preserve_fusions_et_detache_clients_absents(self):
        avant = [("CafeDolbeau", "0015_date_application_gratuite")]
        executor = MigrationExecutor(connection)
        dernier_schema = executor.loader.graph.leaf_nodes()
        try:
            executor.migrate(avant)
            apps = executor.loader.project_state(avant).apps
            AncienClient = apps.get_model("CafeDolbeau", "Client")
            AncienneFusion = apps.get_model("CafeDolbeau", "FusionClient")
            AncienneTransaction = apps.get_model("CafeDolbeau", "TransactionCafe")
            User = apps.get_model("auth", "User")
            employe = User.objects.create(username="migration-fusion-noms")
            source = AncienClient.objects.create(nom_complet="Source actuelle", telephone="migration-source")
            destination = AncienClient.objects.create(nom_complet="Destination actuelle", telephone="migration-destination")
            for source_id, destination_id, statut in (
                (source.pk, destination.pk, "en_attente"),
                (source.pk, destination.pk, "annulee"),
                (999991, destination.pk, "finalisee"),
                (999992, 999993, "finalisee"),
            ):
                AncienneFusion.objects.create(
                    client_source_id=source_id, client_destination_id=destination_id,
                    source_nom="Ancien nom source", destination_nom="Ancien nom destination",
                    cree_par_id=employe.pk, statut=statut,
                )
            AncienneTransaction.objects.create(client=destination, quantite=2, client_origine_id=999991)
            clients = list(AncienClient.objects.order_by("pk").values())
            transactions = list(AncienneTransaction.objects.values())
            fusions = list(AncienneFusion.objects.order_by("pk").values())
            MigrationExecutor(connection).migrate(dernier_schema)
            attendues = []
            for fusion in fusions:
                attendue = {champ: valeur for champ, valeur in fusion.items() if champ not in ("source_nom", "destination_nom")}
                attendue["client_source_id"] = source.pk if fusion["client_source_id"] == source.pk else None
                attendue["client_destination_id"] = destination.pk if fusion["client_destination_id"] == destination.pk else None
                attendues.append(attendue)
            self.assertEqual(list(FusionClient.objects.order_by("pk").values()), attendues)
            self.assertEqual(list(Client.objects.order_by("pk").values()), clients)
            self.assertEqual(list(TransactionCafe.objects.values()), transactions)
            self.assertEqual(str(FusionClient.objects.get(pk=fusions[0]["id"])), "Source actuelle → Destination actuelle")
            connection.check_constraints()
        finally:
            MigrationExecutor(connection).migrate(dernier_schema)

    def test_migration_active_clients_preserve_donnees_sans_fusion(self):
        avant = [("CafeDolbeau", "0013_transactioncafe_client_origine_id")]
        executor = MigrationExecutor(connection)
        dernier_schema = executor.loader.graph.leaf_nodes()
        try:
            executor.migrate(avant)
            apps = executor.loader.project_state(avant).apps
            AncienClient = apps.get_model("CafeDolbeau", "Client")
            AncienneTransaction = apps.get_model("CafeDolbeau", "TransactionCafe")
            AncienneAnnulation = apps.get_model("CafeDolbeau", "AnnulationTransaction")
            User = apps.get_model("auth", "User")
            employe = User.objects.create(username="migration-fusion")
            personne = AncienClient.objects.create(nom_complet="Migration", telephone="111", nombre_cafes_achetes=2)
            fictif = AncienClient.objects.create(pk=-1, nom_complet="Client supprimé")
            ligne = AncienneTransaction.objects.create(client=fictif, quantite=2, employe_id=employe.pk, client_origine_id=123)
            AncienneAnnulation.objects.create(transaction=ligne, employe_id=employe.pk)
            clients = list(AncienClient.objects.values())
            lignes = list(AncienneTransaction.objects.values())
            annulations = list(AncienneAnnulation.objects.values())
            MigrationExecutor(connection).migrate(dernier_schema)
            self.assertEqual(list(Client.objects.values()), [{**client, "est_actif": client["id"] != -1} for client in clients])
            self.assertEqual(list(TransactionCafe.objects.values()), lignes)
            self.assertEqual(list(AnnulationTransaction.objects.values()), annulations)
            self.assertFalse(FusionClient.objects.exists())
            self.assertTrue(Client.objects.get(pk=personne.pk).est_actif)
        finally:
            MigrationExecutor(connection).migrate(dernier_schema)
