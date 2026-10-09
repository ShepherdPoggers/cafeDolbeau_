from datetime import date, datetime
from io import BytesIO
from uuid import uuid4

from django.contrib.auth.models import Group
from django.http import QueryDict
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from ..forms.forms_rapports import RapportForm
from ..models import AnnulationTransaction, Client, Employe, TransactionCafe
from ..rapports import FUSEAU_RAPPORT, generer_rapport


class RapportsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.gerant = Employe.objects.create_user("gerant-rapports")
        cls.gerant.groups.add(Group.objects.get(name="Gérant"))
        cls.barista = Employe.objects.create_user("barista-rapports")
        cls.barista.groups.add(Group.objects.get(name="Barista"))
        cls.personne = Client.objects.create(nom_complet="Camille", telephone="418 555-1111")

    def setUp(self):
        self.client.force_login(self.gerant)

    def instant(self, valeur):
        return datetime.fromisoformat(valeur).replace(tzinfo=FUSEAU_RAPPORT)

    def ligne(self, quantite=1, type_cafe="achat", quand="2026-10-05T10:00:00", auteur=True, annulation=None, operation=None):
        ligne = TransactionCafe.objects.create(
            client=self.personne, quantite=quantite, type_transaction=type_cafe,
            employe=self.barista if auteur else None, operation_id=operation,
        )
        TransactionCafe.objects.filter(pk=ligne.pk).update(date_creation=self.instant(quand))
        if annulation:
            AnnulationTransaction.objects.create(transaction=ligne, employe=self.gerant, date_annulation=self.instant(annulation))
        return ligne

    def filtres(self, **changements):
        donnees = {
            "precision": "periode", "date_debut": "2026-10-01", "date_fin": "2026-10-31",
            "types": ["achat", "prepaye"],
        }
        donnees.update(changements)
        return donnees

    def rapport(self, **changements):
        form = RapportForm(self.filtres(**changements))
        self.assertTrue(form.is_valid(), form.errors)
        return generer_rapport(**form.cleaned_data)

    def classeur(self, **changements):
        response = self.client.get(reverse("rapport_excel"), self.filtres(**changements))
        self.assertEqual(response.status_code, 200)
        contenu = b"".join(response.streaming_content)
        return load_workbook(BytesIO(contenu)), response

    def test_acces_page_et_export_reserves_aux_gerants(self):
        for nom in ("rapports", "rapport_excel"):
            with self.subTest(nom=nom):
                url = reverse(nom)
                self.client.logout()
                self.assertEqual(self.client.get(url).status_code, 302)
                self.client.force_login(self.barista)
                self.assertEqual(self.client.get(url).status_code, 403)
                self.client.force_login(self.gerant)
                self.assertEqual(self.client.get(url).status_code, 200)
                self.assertEqual(self.client.post(url).status_code, 405)

    def test_lien_entete_visible_uniquement_au_gerant(self):
        self.assertContains(self.client.get(reverse("accueil")), 'action="/rapports/"')
        self.client.force_login(self.barista)
        self.assertNotContains(self.client.get(reverse("accueil")), 'action="/rapports/"')

    def test_valeurs_initiales_et_niveau_avant_periode(self):
        response = self.client.get(reverse("rapports"))
        aujourd_hui = timezone.localdate(timezone=FUSEAU_RAPPORT)
        rapport = response.context["rapport"]
        self.assertEqual(rapport.date_debut, aujourd_hui.replace(day=1))
        self.assertEqual(rapport.date_fin, aujourd_hui)
        self.assertEqual(rapport.precision, "periode")
        self.assertEqual(rapport.types, ("achat", "prepaye"))
        self.assertEqual([categorie.code for categorie in rapport.totaux], ["achat", "prepaye"])
        self.assertNotContains(response, "Séparer les achats directs et les prépayés")
        html = response.content.decode()
        self.assertLess(html.index('id="id_precision"'), html.index('id="id_date_debut"'))
        self.assertContains(response, 'type="date"', count=2)

    def test_validation_dates_types_niveau_et_mois_incomplet(self):
        for modifications in (
            {"date_debut": "2026-11-01"}, {"date_fin": "invalide"},
            {"date_fin": "9999-12-31"}, {"types": []}, {"types": ["inconnu"]},
            {"precision": "inconnue"},
            {"precision": "mois", "date_fin": "2026-10-07"},
            {"precision": "mois", "date_debut": "2026-10-02"},
            {"precision": "semaine", "date_debut": "2026-10-06", "date_fin": "2026-10-11"},
            {"precision": "semaine", "date_debut": "2026-10-05", "date_fin": "2026-10-10"},
        ):
            with self.subTest(modifications=modifications):
                form = RapportForm(self.filtres(**modifications))
                self.assertFalse(form.is_valid())
                response = self.client.get(reverse("rapports"), self.filtres(**modifications))
                self.assertNotIn("rapport", response.context)
                self.assertEqual(self.client.get(reverse("rapport_excel"), self.filtres(**modifications)).status_code, 400)

    def test_fevrier_bissextile_et_semaines_entre_annees(self):
        valide = self.rapport(precision="mois", date_debut="2024-02-01", date_fin="2024-02-29")
        self.assertEqual(valide.groupes[0].fin, date(2024, 2, 29))
        self.assertFalse(RapportForm(self.filtres(precision="mois", date_debut="2024-02-01", date_fin="2024-02-28")).is_valid())
        self.assertTrue(RapportForm(self.filtres(precision="mois", date_debut="2025-02-01", date_fin="2025-02-28")).is_valid())
        rapport = self.rapport(precision="semaine", date_debut="2020-12-28", date_fin="2021-01-10")
        self.assertEqual([(g.debut, g.fin) for g in rapport.groupes], [
            (date(2020, 12, 28), date(2021, 1, 3)), (date(2021, 1, 4), date(2021, 1, 10)),
        ])

    def test_achats_et_prepayes_toujours_separes_des_autres_categories(self):
        for type_cafe, quantite in (("achat", 3), ("prepaye", 10), ("gratuit", 2), ("utilise", 4)):
            self.ligne(quantite, type_cafe)
        rapport = self.rapport(types=["achat", "prepaye", "gratuit", "utilise"])
        self.assertEqual([(c.code, c.net) for c in rapport.totaux], [("achat", 3), ("prepaye", 10), ("gratuit", 2), ("utilise", 4)])
        rapport = self.rapport()
        self.assertEqual([(c.code, c.net) for c in rapport.totaux], [("achat", 3), ("prepaye", 10)])

    def test_annulation_dans_meme_periode_et_auteurs_evenements(self):
        ligne = self.ligne(3, annulation="2026-10-06T11:00:00")
        rapport = self.rapport(precision="transaction")
        total = rapport.totaux[0]
        self.assertEqual((total.enregistrees, total.annulees, total.net), (3, 3, 0))
        self.assertEqual([(e.transaction_id, e.nature, e.employe, e.net) for e in rapport.evenements], [
            (ligne.pk, "Transaction", "barista-rapports", 3), (ligne.pk, "Annulation", "gerant-rapports", -3),
        ])

    def test_annulation_autre_mois_ne_change_pas_les_ventes_du_mois_original(self):
        self.ligne(10, quand="2026-09-30T23:30:00", annulation="2026-10-02T09:00:00")
        septembre = self.rapport(date_debut="2026-09-01", date_fin="2026-09-30")
        octobre = self.rapport()
        self.assertEqual((septembre.totaux[0].enregistrees, septembre.totaux[0].annulees, septembre.totaux[0].net), (10, 0, 10))
        self.assertEqual((octobre.totaux[0].enregistrees, octobre.totaux[0].annulees, octobre.totaux[0].net), (0, 10, -10))
        self.assertEqual(self.rapport(date_debut="2026-09-01").totaux[0].net, 0)

    def test_filtre_type_sapplique_aux_transactions_et_annulations(self):
        self.ligne(10, "prepaye", annulation="2026-10-06T12:00:00")
        self.ligne(2, "achat")
        rapport = self.rapport(types=["achat"], precision="transaction")
        self.assertEqual(rapport.totaux[0].net, 2)
        self.assertEqual(len(rapport.evenements), 1)

    def test_operation_multiligne_comptee_par_type_et_ancienne_inconnue(self):
        operation = uuid4()
        self.ligne(1, "achat", operation=operation, annulation="2026-10-07T10:00:00")
        self.ligne(1, "gratuit", operation=operation, annulation="2026-10-07T10:00:00")
        ancienne = self.ligne(2, auteur=False)
        rapport = self.rapport(precision="transaction", types=["achat", "gratuit"])
        self.assertEqual([(c.enregistrees, c.annulees, c.net) for c in rapport.totaux], [(3, 1, 2), (1, 1, 0)])
        self.assertEqual(next(e.employe for e in rapport.evenements if e.transaction_id == ancienne.pk), "Inconnu")

    def test_bornes_inclusives_a_toronto_et_regroupement_journalier(self):
        self.ligne(1, quand="2026-10-04T23:59:59")
        self.ligne(2, quand="2026-10-05T00:00:00")
        self.ligne(3, quand="2026-10-05T23:59:59.999999")
        self.ligne(4, quand="2026-10-06T00:00:00")
        with timezone.override("UTC"):
            rapport = self.rapport(date_debut="2026-10-05", date_fin="2026-10-05", precision="jour")
        self.assertEqual(rapport.totaux[0].net, 5)
        self.assertEqual(rapport.groupes[0].categories[0].net, 5)

    def test_changements_heure_jours_de_23_et_25_heures(self):
        for debut, heure in (("2026-03-08", "23:30:00"), ("2026-11-01", "23:30:00")):
            with self.subTest(debut=debut):
                ligne = self.ligne(1, quand=f"{debut}T{heure}")
                rapport = self.rapport(date_debut=debut, date_fin=debut, precision="transaction")
                self.assertEqual(rapport.totaux[0].net, 1)
                self.assertEqual(rapport.evenements[0].date.hour, 23)
                self.assertEqual(rapport.evenements[0].transaction_id, ligne.pk)
        ligne = self.ligne(quand="2026-11-01T01:30:00")
        deuxieme = self.instant("2026-11-01T01:30:00").replace(fold=1)
        TransactionCafe.objects.filter(pk=ligne.pk).update(date_creation=deuxieme)
        rapport = self.rapport(date_debut="2026-11-01", date_fin="2026-11-01", precision="transaction")
        self.assertEqual(rapport.totaux[0].net, 2)

    def test_periodes_sans_activite_et_regroupements_chronologiques(self):
        self.ligne(3, quand="2026-10-02T10:00:00")
        self.ligne(4, quand="2026-10-03T10:00:00")
        rapport = self.rapport(precision="jour", date_fin="2026-10-04")
        self.assertEqual([g.categories[0].net for g in rapport.groupes], [0, 3, 4, 0])
        self.assertEqual(self.rapport(precision="mois", date_fin="2026-12-31").groupes[2].categories[0].net, 0)
        rapport = self.rapport(precision="semaine", date_debut="2026-09-28", date_fin="2026-10-11")
        self.assertEqual([(g.debut, g.fin, g.categories[0].net) for g in rapport.groupes], [
            (date(2026, 9, 28), date(2026, 10, 4), 7), (date(2026, 10, 5), date(2026, 10, 11), 0),
        ])

    def test_ordre_reel_pendant_heure_repetee_et_affichage_toujours_a_toronto(self):
        premiere = self.ligne(quand="2026-11-01T01:30:00")
        seconde = self.ligne(quand="2026-11-01T01:15:00")
        TransactionCafe.objects.filter(pk=seconde.pk).update(date_creation=self.instant("2026-11-01T01:15:00").replace(fold=1))
        rapport = self.rapport(precision="transaction", date_debut="2026-11-01", date_fin="2026-11-01")
        self.assertEqual([e.transaction_id for e in rapport.evenements], [premiere.pk, seconde.pk])
        self.ligne(quand="2026-10-05T23:30:00")
        with timezone.override("UTC"):
            response = self.client.get(reverse("rapports"), self.filtres(precision="transaction", date_debut="2026-10-05", date_fin="2026-10-05"))
        self.assertContains(response, "05/10/2026 23:30:00")

    def test_excel_mensuel_concorde_avec_les_groupes_et_totaux_affiches(self):
        self.ligne(10, "prepaye", quand="2026-09-30T23:00:00", annulation="2026-10-02T10:00:00")
        filtres = {"precision": "mois", "date_debut": "2026-09-01", "date_fin": "2026-11-30", "types": ["prepaye"]}
        response = self.client.get(reverse("rapports"), self.filtres(**filtres))
        rapport = response.context["rapport"]
        classeur, _ = self.classeur(**filtres)
        detail = classeur["Détail des totaux"]
        lignes = list(detail.iter_rows(min_row=2, values_only=True))
        for groupe, ligne in zip(rapport.groupes, lignes[1::2]):
            categorie = groupe.categories[0]
            self.assertEqual(ligne, (categorie.libelle, categorie.enregistrees, categorie.annulees, categorie.net))
        self.assertEqual([ligne[3] for ligne in lignes[1::2]], [10, -10, 0, 0])
        self.assertEqual([ligne[0] for ligne in lignes[::2]], ["2026-09", "2026-10", "2026-11", "TOTAL"])
        matrice = list(classeur["Synthèse"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(matrice, [("2026-09", 10, 10), ("2026-10", -10, -10), ("2026-11", 0, 0), ("TOTAL", 0, 0)])

    def test_lecture_jointe_unique_et_aucune_modification_des_donnees(self):
        self.ligne(3, annulation="2026-10-07T10:00:00")
        form = RapportForm(self.filtres(precision="transaction"))
        self.assertTrue(form.is_valid())
        avant = list(Client.objects.values())
        transactions = list(TransactionCafe.objects.values())
        traces = list(AnnulationTransaction.objects.values())
        with self.assertNumQueries(1):
            generer_rapport(**form.cleaned_data)
        self.classeur(precision="transaction")
        self.assertEqual(avant, list(Client.objects.values()))
        self.assertEqual(transactions, list(TransactionCafe.objects.values()))
        self.assertEqual(traces, list(AnnulationTransaction.objects.values()))

    def test_pagination_garde_filtres_et_totaux_et_exporte_toutes_les_lignes(self):
        for _ in range(55):
            self.ligne()
        donnees = self.filtres(precision="transaction")
        response = self.client.get(reverse("rapports"), donnees)
        self.assertEqual(len(response.context["page_evenements"]), 50)
        self.assertEqual(response.context["rapport"].totaux[0].net, 55)
        filtres = QueryDict(response.context["filtres_pagination"])
        self.assertEqual(filtres.getlist("types"), ["achat", "prepaye"])
        response = self.client.get(reverse("rapports"), {**donnees, "page": 2})
        self.assertEqual(len(response.context["page_evenements"]), 5)
        self.assertEqual(response.context["rapport"].totaux[0].net, 55)
        classeur, _ = self.classeur(precision="transaction", page=2)
        self.assertEqual(classeur["Transactions"].max_row, 56)

    def test_classeur_dates_nombres_annulations_parametres_et_nom(self):
        self.ligne(3, annulation="2026-10-06T11:30:00")
        classeur, response = self.classeur(precision="transaction")
        self.assertEqual(classeur.sheetnames, ["Paramètres", "Synthèse", "Détail des totaux", "Transactions"])
        self.assertIn("rapport_cafes_2026-10-01_2026-10-31.xlsx", response["Content-Disposition"])
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("no-store", response["Cache-Control"])
        synthese = classeur["Détail des totaux"]
        self.assertEqual([synthese.cell(3, colonne).value for colonne in (2, 3, 4)], [3, 3, 0])
        self.assertEqual(synthese["A2"].value, "2026-10-01 au 2026-10-31")
        self.assertEqual(synthese["D3"].data_type, "n")
        self.assertEqual(synthese.freeze_panes, "A2")
        self.assertEqual(classeur["Synthèse"].freeze_panes, "B2")
        self.assertEqual(classeur["Synthèse"]["B2"].value, 0)
        self.assertEqual(classeur["Synthèse"]["B2"].data_type, "n")
        self.assertEqual(classeur["Transactions"]["A3"].value, datetime(2026, 10, 6, 11, 30))
        self.assertEqual(classeur["Transactions"]["I3"].value, -3)
        parametres = {ligne[0]: ligne[1] for ligne in classeur["Paramètres"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(parametres["Fuseau horaire"], "America/Toronto")
        self.assertEqual(parametres["Niveau de détail"], "Par transaction")

    def test_classeur_synthese_seule_et_noms_traites_comme_texte(self):
        Client.objects.filter(pk=self.personne.pk).update(nom_complet="=1+2")
        self.ligne(auteur=False)
        classeur, _ = self.classeur()
        self.assertEqual(classeur.sheetnames, ["Paramètres", "Synthèse", "Détail des totaux"])
        classeur, _ = self.classeur(precision="transaction")
        self.assertEqual(classeur["Transactions"]["E2"].value, "=1+2")
        self.assertEqual(classeur["Transactions"]["E2"].data_type, "s")
        self.assertEqual(classeur["Transactions"]["F2"].value, "Inconnu")

    def test_page_detail_vide_et_export_vide(self):
        response = self.client.get(reverse("rapports"), self.filtres(precision="transaction"))
        self.assertContains(response, "Aucune transaction ni annulation")
        classeur, _ = self.classeur(precision="transaction")
        self.assertEqual(classeur["Transactions"].max_row, 1)
        self.assertEqual(classeur["Synthèse"]["B2"].value, 0)

    def test_tableau_categories_en_colonnes_et_totaux_par_ligne_et_colonne(self):
        self.ligne(3, "achat", quand="2026-10-01T10:00:00")
        self.ligne(10, "prepaye", quand="2026-10-01T10:00:00", annulation="2026-10-02T10:00:00")
        self.ligne(2, "gratuit", quand="2026-10-01T10:00:00")
        self.ligne(4, "utilise", quand="2026-10-02T10:00:00")
        filtres = self.filtres(precision="jour", date_fin="2026-10-03", types=["achat", "prepaye", "gratuit", "utilise"])
        response = self.client.get(reverse("rapports"), filtres)
        self.assertContains(response, '<tr><th scope="row">2026-10-01</th><td>3</td><td>10</td><td>2</td><td>0</td><td class="rapport-total">15</td></tr>', html=True)
        self.assertContains(response, '<tr><th scope="row">2026-10-02</th><td>0</td><td>-10</td><td>0</td><td>4</td><td class="rapport-total">-6</td></tr>', html=True)
        self.assertContains(response, '<tr><th scope="row">TOTAL</th><td>3</td><td>0</td><td>2</td><td>4</td><td>9</td></tr>', html=True)
        classeur, _ = self.classeur(**filtres)
        lignes = list(classeur["Synthèse"].iter_rows(values_only=True))
        self.assertEqual(lignes[0], ("Date", "Achats directs", "Cafés prépayés", "Cafés gratuits accordés", "Cafés consommés (prépayés ou gratuits)", "TOTAL"))
        self.assertEqual(lignes[1][1:], (3, 10, 2, 0, 15))
        self.assertEqual(lignes[2][1:], (0, -10, 0, 4, -6))
        self.assertEqual(lignes[3][1:], (0, 0, 0, 0, 0))
        self.assertEqual(lignes[4], ("TOTAL", 3, 0, 2, 4, 9))

    def test_tableau_separe_les_ventes_et_ne_montre_que_les_categories_selectionnees(self):
        self.ligne(2, "achat")
        self.ligne(10, "prepaye")
        self.ligne(5, "utilise")
        classeur, _ = self.classeur(precision="jour", date_debut="2026-10-05", date_fin="2026-10-05")
        lignes = list(classeur["Synthèse"].iter_rows(values_only=True))
        self.assertEqual(lignes[0], ("Date", "Achats directs", "Cafés prépayés", "TOTAL"))
        self.assertEqual(lignes[1][1:], (2, 10, 12))
        self.assertEqual(lignes[2], ("TOTAL", 2, 10, 12))
        classeur, _ = self.classeur(precision="jour", date_debut="2026-10-05", date_fin="2026-10-05", separer_ventes="false")
        self.assertEqual(tuple(cellule.value for cellule in classeur["Synthèse"][1]), lignes[0])
        self.assertNotIn("Séparer les ventes", [ligne[0] for ligne in classeur["Paramètres"].iter_rows(values_only=True)])

    def test_detail_en_blocs_par_date_et_categories_avec_annulations_et_jours_vides(self):
        self.ligne(3, "achat", quand="2026-10-01T10:00:00", annulation="2026-10-02T11:00:00")
        self.ligne(10, "prepaye", quand="2026-10-02T12:00:00")
        filtres = self.filtres(precision="jour", date_fin="2026-10-03")
        response = self.client.get(reverse("rapports"), filtres)
        self.assertContains(response, '<tr class="rapport-detail-periode"><th colspan="4" scope="rowgroup">2026-10-01</th></tr>', html=True)
        self.assertContains(response, '<tr><th scope="row">Achats directs</th><td>0</td><td>3</td><td>-3</td></tr>', html=True)
        self.assertContains(response, '<tr><th scope="row">Cafés prépayés</th><td>10</td><td>0</td><td>10</td></tr>', html=True)
        self.assertContains(response, 'class="rapport-detail-groupe rapport-detail-clair"', count=2)
        self.assertContains(response, 'class="rapport-detail-groupe rapport-detail-fonce"', count=1)
        classeur, _ = self.classeur(**filtres)
        detail = classeur["Détail des totaux"]
        self.assertEqual(tuple(c.value for c in detail[1]), ("Catégorie", "Enregistrées", "Annulées", "Net"))
        self.assertEqual(str(detail.merged_cells), "A2:D2 A5:D5 A8:D8 A11:D11")
        self.assertEqual(detail["A2"].value, datetime(2026, 10, 1))
        self.assertEqual(detail["A2"].number_format, "yyyy-mm-dd")
        self.assertEqual(tuple(c.value for c in detail[6]), ("Achats directs", 0, 3, -3))
        self.assertEqual(tuple(c.value for c in detail[7]), ("Cafés prépayés", 10, 0, 10))
        self.assertEqual(tuple(c.value for c in detail[9]), ("Achats directs", 0, 0, 0))
        self.assertEqual(tuple(c.value for c in detail[12]), ("Achats directs", 3, 3, 0))
        self.assertEqual(tuple(c.value for c in detail[13]), ("Cafés prépayés", 10, 0, 10))
        self.assertEqual(detail["B3"].fill.fgColor.rgb, "0078A3CF")
        self.assertEqual(detail["B6"].fill.fgColor.rgb, "003768A0")
        self.assertTrue(detail["A2"].font.bold)
        self.assertEqual(detail["A2"].alignment.horizontal, "center")
