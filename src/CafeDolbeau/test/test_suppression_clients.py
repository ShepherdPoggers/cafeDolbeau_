from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict
from threading import Barrier
from unittest.mock import patch

from django.contrib.admin.models import DELETION, LogEntry
from django.core import signing
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import close_old_connections, connection, transaction
from django.db.migrations.executor import MigrationExecutor
from django.db.models.signals import pre_delete
from django.test import Client as Navigateur, TestCase, TransactionTestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from ..forms.forms_annulation import SelectionAnciennesForm
from ..models import AnnulationTransaction, Client, Employe, TransactionCafe
from ..permissions import peut_annuler
from ..rapports import exporter_excel, generer_rapport
from ..services import ajouter_cafes, annuler_operation, supprimer_clients
from .test_annulations import creer_employe


class SuppressionClientTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.gerant = creer_employe("gerant-suppression", "Gérant")
        cls.barista = creer_employe("barista-suppression")
        cls.personne = Client.objects.create(
            nom_complet="Camille", telephone="418 555-1000", nombre_cafes_achetes=9,
            nombre_cafes_prepayes=4, nombre_cafes_gratuits=2,
        )
        cls.autre = Client.objects.create(nom_complet="Alex", courriel="alex@example.com")

    def setUp(self):
        self.client.force_login(self.gerant)
        self.url = reverse("admin:CafeDolbeau_client_delete", args=[self.personne.pk])
        self.liste = reverse("admin:CafeDolbeau_client_changelist")

    def operation(self, personne=None, quantite=1, type_cafe="achat"):
        return ajouter_cafes(personne or self.personne, quantite, type_cafe, employe=self.barista)

    def confirmer(self, url=None):
        url = url or self.url
        apercu = self.client.get(url)
        self.assertEqual(apercu.status_code, 200)
        return self.client.post(url, {"post": "yes", "confirmation": apercu.context["confirmation"]})

    def test_transfert_conserve_tous_les_champs_et_annulations(self):
        lignes = self.operation()
        annuler_operation(lignes[0].pk, self.gerant)
        avant = list(TransactionCafe.objects.order_by("pk").values())
        audits = list(AnnulationTransaction.objects.values())
        self.assertEqual(supprimer_clients([self.personne.pk], employe=self.gerant), 1)
        self.assertFalse(Client.objects.filter(pk=self.personne.pk).exists())
        apres = list(TransactionCafe.objects.order_by("pk").values())
        for original, transfere in zip(avant, apres):
            self.assertEqual(transfere, {
                **original, "client_id": -1, "client_origine_id": self.personne.pk,
            })
        self.assertEqual(AnnulationTransaction.objects.count(), len(audits))
        self.assertEqual(list(AnnulationTransaction.objects.values()), audits)
        fictif = Client.objects.get(pk=-1)
        fictif.full_clean()
        self.assertEqual((fictif.nom_complet, fictif.telephone, fictif.courriel), ("Client supprimé", "", ""))
        self.assertEqual((fictif.nombre_cafes_achetes, fictif.nombre_cafes_prepayes, fictif.nombre_cafes_gratuits), (0, 0, 0))
        trace = LogEntry.objects.get(action_flag=DELETION)
        self.assertEqual((trace.user_id, trace.object_id), (self.gerant.pk, str(self.personne.pk)))

    def test_sans_transactions_et_reutilisation_du_client_fictif(self):
        self.assertEqual(supprimer_clients([self.personne.pk], employe=self.gerant), 1)
        self.assertEqual(supprimer_clients([self.autre.pk], employe=self.gerant), 1)
        self.assertEqual(list(Client.objects.values_list("pk", flat=True)), [-1])
        self.assertEqual(supprimer_clients([self.personne.pk], employe=self.gerant), 0)
        self.assertEqual(LogEntry.objects.filter(action_flag=DELETION).count(), 2)

    def test_suppression_en_lot_et_origines_distinctes(self):
        ligne = self.operation()[0]
        autre = self.operation(self.autre)[0]
        self.assertEqual(supprimer_clients([self.personne.pk, self.autre.pk], employe=self.gerant), 2)
        for source, origine in ((ligne, self.personne.pk), (autre, self.autre.pk)):
            source.refresh_from_db()
            self.assertEqual((source.client_id, source.client_origine_id), (-1, origine))

    def test_service_reverifie_les_permissions_et_refuse_le_barista(self):
        with self.assertRaises(PermissionDenied):
            supprimer_clients([self.personne.pk], employe=self.barista)
        self.assertTrue(self.gerant.has_perm("CafeDolbeau.acces_gerant"))
        self.gerant.groups.clear()
        with self.assertRaises(PermissionDenied):
            supprimer_clients([self.personne.pk], employe=self.gerant)
        self.assertTrue(Client.objects.filter(pk=self.personne.pk).exists())
        self.assertFalse(Client.objects.filter(pk=-1).exists())
        self.assertFalse(LogEntry.objects.exists())

    def test_client_fictif_protege_et_pas_de_nouvel_achat(self):
        supprimer_clients([self.personne.pk], employe=self.gerant)
        fictif = Client.objects.get(pk=-1)
        for supprimer in (fictif.delete, lambda: Client.objects.filter(pk=-1).delete()):
            with self.assertRaises(ValidationError), transaction.atomic():
                supprimer()
        with self.assertRaises(ValidationError):
            supprimer_clients([-1, self.autre.pk], employe=self.gerant)
        with self.assertRaises(ValidationError):
            ajouter_cafes(fictif, 1, employe=self.gerant)
        self.assertTrue(Client.objects.filter(pk=self.autre.pk).exists())
        self.assertFalse(TransactionCafe.objects.exists())
        with self.assertRaises(ValidationError):
            Client(nom_complet="Sans coordonnées").full_clean()

    def test_identifiants_positifs_non_reutilises(self):
        dernier = Client.objects.create(nom_complet="Dernier", courriel="dernier@example.com")
        supprimer_clients([dernier.pk], employe=self.gerant)
        suivant = Client.objects.create(nom_complet="Suivant", courriel="suivant@example.com")
        self.assertGreater(suivant.pk, dernier.pk)

    def test_echec_sur_second_client_annule_transfert_journal_et_suppression(self):
        self.operation()
        self.operation(self.autre)
        clients = list(Client.objects.values())
        transactions = list(TransactionCafe.objects.values())

        def echouer(sender, instance, **kwargs):
            if instance.pk == self.autre.pk:
                raise RuntimeError("échec suppression")

        pre_delete.connect(echouer, sender=Client)
        try:
            with self.assertRaises(RuntimeError):
                supprimer_clients([self.personne.pk, self.autre.pk], employe=self.gerant)
        finally:
            pre_delete.disconnect(echouer, sender=Client)
        self.assertEqual(list(Client.objects.values()), clients)
        self.assertEqual(list(TransactionCafe.objects.values()), transactions)
        self.assertFalse(LogEntry.objects.exists())

    def test_echec_journal_annule_transfert_et_creation_fictive(self):
        ligne = self.operation()[0]
        with patch.object(LogEntry.objects, "log_actions", side_effect=RuntimeError("échec journal")):
            with self.assertRaises(RuntimeError):
                supprimer_clients([self.personne.pk], employe=self.gerant)
        ligne.refresh_from_db()
        self.assertEqual(ligne.client_id, self.personne.pk)
        self.assertIsNone(ligne.client_origine_id)
        self.assertFalse(Client.objects.filter(pk=-1).exists())

    def test_confirmation_individuelle_sans_ecriture_et_avertissement(self):
        self.operation()
        clients = list(Client.objects.values())
        transactions = list(TransactionCafe.objects.values())
        response = self.client.get(self.url)
        self.assertContains(response, "Cafés prépayés perdus")
        self.assertContains(response, "Cafés gratuits perdus")
        self.assertContains(response, "seront conservées")
        self.assertEqual(response.context["clients"][0].nombre_cafes_prepayes, 4)
        self.assertEqual(list(Client.objects.values()), clients)
        self.assertEqual(list(TransactionCafe.objects.values()), transactions)
        self.assertFalse(LogEntry.objects.exists())
        self.assertRedirects(self.confirmer(), self.liste)
        self.assertEqual(LogEntry.objects.filter(action_flag=DELETION).count(), 1)

    def test_confirmation_en_lot_sans_ecriture_puis_suppression(self):
        self.operation()
        self.operation(self.autre)
        donnees = {"action": "delete_selected", "_selected_action": [self.personne.pk, self.autre.pk]}
        response = self.client.post(self.liste, donnees)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["clients"]), 2)
        self.assertFalse(Client.objects.filter(pk=-1).exists())
        response = self.client.post(self.liste, {
            **donnees, "post": "yes", "confirmation": response.context["confirmation"],
        })
        self.assertRedirects(response, self.liste)
        self.assertEqual(Client.objects.count(), 1)
        self.assertEqual(TransactionCafe.objects.filter(client_id=-1).count(), 3)
        self.assertEqual(LogEntry.objects.filter(action_flag=DELETION).count(), 2)

    def test_confirmation_perimee_reaffiche_soldes_et_ne_supprime_pas(self):
        jeton = self.client.get(self.url).context["confirmation"]
        self.operation(quantite=1, type_cafe="prepaye")
        response = self.client.post(self.url, {"post": "yes", "confirmation": jeton})
        self.assertContains(response, "soldes ont changé", status_code=409)
        self.assertEqual(response.context["clients"][0].nombre_cafes_prepayes, 14)
        self.assertTrue(Client.objects.filter(pk=self.personne.pk).exists())
        self.assertFalse(LogEntry.objects.exists())
        response = self.client.post(self.url, {"post": "yes", "confirmation": response.context["confirmation"]})
        self.assertRedirects(response, self.liste)

    def test_soldes_reverifies_sous_verrou_meme_apres_lecture_admin(self):
        jeton = self.client.get(self.url).context["confirmation"]

        def modifier_avant_service(ids, **kwargs):
            Client.objects.filter(pk=self.personne.pk).update(nombre_cafes_prepayes=7)
            return supprimer_clients(ids, **kwargs)

        with patch("CafeDolbeau.admin.supprimer_clients", side_effect=modifier_avant_service):
            response = self.client.post(self.url, {"post": "yes", "confirmation": jeton})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.context["clients"][0].nombre_cafes_prepayes, 7)
        self.assertFalse(LogEntry.objects.exists())

    def test_changement_soldes_en_lot_bloque_toute_la_suppression(self):
        donnees = {"action": "delete_selected", "_selected_action": [self.personne.pk, self.autre.pk]}
        jeton = self.client.post(self.liste, donnees).context["confirmation"]
        self.operation(self.autre)
        response = self.client.post(self.liste, {**donnees, "post": "yes", "confirmation": jeton})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(Client.objects.count(), 2)
        self.assertFalse(LogEntry.objects.exists())

    def test_confirmation_invalide_expiree_ou_pour_un_autre_client(self):
        jeton = self.client.get(self.url).context["confirmation"]
        with patch("django.core.signing.time.time", return_value=0):
            expire = signing.dumps({"employe": self.gerant.pk, "soldes": {}}, salt="suppression-client")
        for confirmation in ("", "falsifie", expire):
            with self.subTest(confirmation=confirmation):
                response = self.client.post(self.url, {"post": "yes", "confirmation": confirmation})
                self.assertEqual(response.status_code, 400)
        autre_url = reverse("admin:CafeDolbeau_client_delete", args=[self.autre.pk])
        self.assertEqual(self.client.post(autre_url, {"post": "yes", "confirmation": jeton}).status_code, 409)
        self.assertEqual(Client.objects.count(), 2)
        self.assertFalse(LogEntry.objects.exists())

    def test_confirmation_non_transferable_entre_employes(self):
        jeton = self.client.get(self.url).context["confirmation"]
        autre_gerant = creer_employe("autre-gerant", "Gérant")
        self.client.force_login(autre_gerant)
        self.assertEqual(self.client.post(self.url, {"post": "yes", "confirmation": jeton}).status_code, 400)
        self.assertTrue(Client.objects.filter(pk=self.personne.pk).exists())

    def test_achat_http_client_supprime_apres_chargement_retourne_404(self):
        def supprimer_avant_achat(client, quantite, type_transaction, *, employe):
            supprimer_clients([client.pk], employe=self.gerant)
            return ajouter_cafes(client, quantite, type_transaction, employe=employe)

        with patch("CafeDolbeau.views.views.ajouter_cafes", side_effect=supprimer_avant_achat):
            response = self.client.post(reverse("ajouter_cafes", args=[self.personne.pk]), {
                "quantite": 1, "type_transaction": "achat",
            }, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.status_code, 404)
        self.assertFalse(TransactionCafe.objects.exists())

    def test_confirmation_annulation_puis_suppression_redirige_vers_admin(self):
        ligne = self.operation()[0]
        url = reverse("annuler_transaction", args=[self.personne.pk, ligne.pk])
        jeton = self.client.get(url).context["confirmation"]

        def supprimer_avant_annulation(transaction_id, employe, details=None):
            supprimer_clients([self.personne.pk], employe=self.gerant)
            return annuler_operation(transaction_id, employe, details)

        with patch("CafeDolbeau.views.views.annuler_operation", side_effect=supprimer_avant_annulation):
            response = self.client.post(url, {"action": "confirmer", "confirmation": jeton})
        self.assertRedirects(response, reverse("admin:CafeDolbeau_transactioncafe_changelist"))
        self.assertTrue(TransactionCafe.objects.get(pk=ligne.pk).annulee)

    def test_acces_http_anonyme_barista_sans_role_inactif_et_superutilisateur(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
        self.assertEqual(self.client.post(self.url, {"post": "yes"}).status_code, 302)
        sans_role = Employe.objects.create_user("sans-role-suppression")
        for employe in (self.barista, sans_role):
            self.client.force_login(employe)
            self.assertEqual(self.client.get(self.url, follow=True).status_code, 403)
            self.assertEqual(self.client.post(self.url, {"post": "yes"}, follow=True).status_code, 403)
            self.assertEqual(self.client.post(self.liste, {
                "action": "delete_selected", "_selected_action": [self.personne.pk], "post": "yes",
            }, follow=True).status_code, 403)
        self.gerant.is_active = False
        self.gerant.save(update_fields=["is_active"])
        with self.assertRaises(PermissionDenied):
            supprimer_clients([self.personne.pk], employe=self.gerant)
        self.assertTrue(Client.objects.filter(pk=self.personne.pk).exists())
        superutilisateur = Employe.objects.create_superuser("super-suppression")
        self.client.force_login(superutilisateur)
        self.assertRedirects(self.confirmer(), self.liste)

    def test_csrf_individuel_et_en_lot(self):
        navigateur = Navigateur(enforce_csrf_checks=True)
        navigateur.force_login(self.gerant)
        jeton = navigateur.get(self.url).context["confirmation"]
        self.assertEqual(navigateur.post(self.url, {"post": "yes", "confirmation": jeton}).status_code, 403)
        self.assertEqual(navigateur.post(self.liste, {
            "action": "delete_selected", "_selected_action": [self.personne.pk],
        }).status_code, 403)
        self.assertEqual(navigateur.post(self.url, {
            "post": "yes", "confirmation": jeton,
            "csrfmiddlewaretoken": navigateur.cookies["csrftoken"].value,
        }).status_code, 302)

    def test_fictif_masque_dans_recherches_et_gestion_clients(self):
        supprimer_clients([self.personne.pk], employe=self.gerant)
        for parametres in ({}, {"q": "Client supprimé"}, {"q": "Client supprimé", "partiel": "recherche"}):
            response = self.client.get(reverse("accueil"), parametres)
            self.assertNotContains(response, '<td><span class="cell-text" title="Client supprimé">')
            self.assertFalse(response.context["resultats"].filter(pk=-1).exists())
        self.assertEqual(self.client.get(self.liste).context["cl"].result_count, 1)
        for action in ("change", "delete"):
            url = reverse(f"admin:CafeDolbeau_client_{action}", args=[-1])
            self.assertIn(self.client.get(url).status_code, (302, 404))
            self.assertIn(self.client.post(url, {"post": "yes"}).status_code, (302, 404))
        self.assertTrue(Client.objects.filter(pk=-1).exists())

    def test_annulation_multiligne_apres_suppression_sans_crediter_fictif(self):
        # Une consommation restaure normalement des crédits : aucun ici.
        lignes = self.operation(quantite=10, type_cafe="utilise")
        supprimer_clients([self.personne.pk], employe=self.gerant)
        source = TransactionCafe.objects.get(pk=lignes[0].pk)
        self.assertFalse(peut_annuler(source, self.barista))
        with self.assertRaises(PermissionDenied):
            annuler_operation(source.pk, self.barista)
        url = reverse("annuler_transaction_admin", args=[source.pk])
        self.client.force_login(self.barista)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.assertEqual(self.client.post(url, {"action": "confirmer"}).status_code, 403)
        self.client.force_login(self.gerant)
        apercu = self.client.get(url)
        self.assertContains(apercu, "sans modifier de solde")
        self.assertEqual(apercu.context["soldes"], [])
        confirmation = {"action": "confirmer", "confirmation": apercu.context["confirmation"]}
        for _ in range(2):
            self.assertRedirects(self.client.post(url, confirmation), reverse("admin:CafeDolbeau_transactioncafe_changelist"))
        self.assertEqual(AnnulationTransaction.objects.count(), len(lignes))
        fictif = Client.objects.get(pk=-1)
        self.assertEqual((fictif.nombre_cafes_achetes, fictif.nombre_cafes_prepayes, fictif.nombre_cafes_gratuits), (0, 0, 0))
        ancien_url = reverse("annuler_transaction", args=[self.personne.pk, source.pk])
        self.assertEqual(self.client.get(ancien_url).status_code, 404)

    def test_anciennes_transactions_ne_melangent_pas_deux_origines(self):
        achat = TransactionCafe.objects.create(client=self.personne, quantite=2)
        consommation = TransactionCafe.objects.create(client=self.personne, quantite=3, type_transaction="utilise")
        autre = TransactionCafe.objects.create(client=self.autre, quantite=1)
        supprimer_clients([self.personne.pk, self.autre.pk], employe=self.gerant)
        achat.refresh_from_db()
        formulaire = SelectionAnciennesForm(client=achat.client, source=achat)
        self.assertEqual(set(formulaire.fields["lignes"].queryset.values_list("pk", flat=True)), {achat.pk, consommation.pk})
        url = reverse("annuler_transaction_admin", args=[achat.pk])
        response = self.client.post(url, {"action": "completer", "lignes": [achat.pk, autre.pk]})
        self.assertTrue(response.context["selection"].errors)
        with self.assertRaises(ValidationError):
            annuler_operation(achat.pk, self.gerant, {achat.pk: (2, 0, 0), autre.pk: (1, 0, 0)})
        valeurs = {achat.pk: (2, 0, 0), consommation.pk: (0, -1, -2)}
        donnees = {
            "action": "previsualiser", "lignes": list(valeurs), "form-TOTAL_FORMS": 2,
            "form-INITIAL_FORMS": 2, "form-MIN_NUM_FORMS": 0, "form-MAX_NUM_FORMS": 2000,
        }
        for index, (pk, variations) in enumerate(valeurs.items()):
            donnees[f"form-{index}-transaction_id"] = pk
            for champ, valeur in zip(("variation_achetes", "variation_prepayes", "variation_gratuits"), variations):
                donnees[f"form-{index}-{champ}"] = valeur
        response = self.client.post(url, donnees)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.client.post(url, {
            "action": "confirmer", "confirmation": response.context["confirmation"],
        }).status_code, 302)
        achat.refresh_from_db()
        consommation.refresh_from_db()
        self.assertEqual(achat.operation_id, consommation.operation_id)
        self.assertEqual(AnnulationTransaction.objects.count(), 2)
        self.assertFalse(TransactionCafe.objects.get(pk=autre.pk).annulee)
        self.assertEqual(Client.objects.get(pk=-1).nombre_cafes_prepayes, 0)

    def test_admin_transactions_transferees_et_rapports_excel_conserves(self):
        self.operation()
        self.operation(quantite=1, type_cafe="prepaye")
        annuler_operation(TransactionCafe.objects.first().pk, self.gerant)
        jour = timezone.localdate()

        def rapport():
            return generer_rapport(date_debut=jour, date_fin=jour, types=tuple(TransactionCafe.Type.values), precision="transaction")

        avant = rapport()
        supprimer_clients([self.personne.pk], employe=self.gerant)
        apres = rapport()
        self.assertEqual([asdict(total) for total in avant.totaux], [asdict(total) for total in apres.totaux])
        self.assertEqual(len(avant.evenements), len(apres.evenements))
        for original, transfere in zip(avant.evenements, apres.evenements):
            self.assertEqual(asdict(transfere), {**asdict(original), "client": "Client supprimé"})
        classeur = load_workbook(exporter_excel(apres))
        self.assertEqual({ligne[4] for ligne in classeur["Transactions"].iter_rows(min_row=2, values_only=True)}, {"Client supprimé"})
        url = reverse("admin:CafeDolbeau_transactioncafe_changelist")
        self.assertContains(self.client.get(url), "Client supprimé")
        ligne = TransactionCafe.objects.filter(annulation__isnull=True).first()
        self.assertContains(self.client.get(url), reverse("annuler_transaction_admin", args=[ligne.pk]))


class SuppressionConcurrenceTests(TransactionTestCase):
    def setUp(self):
        self.gerant = creer_employe("gerant-concurrence-suppression", "Gérant")
        self.barista = creer_employe("barista-concurrence-suppression")
        self.personne = Client.objects.create(nom_complet="Camille", telephone="418 555-1000")

    def simultanement(self, fonctions):
        barriere = Barrier(len(fonctions))

        def executer(fonction):
            close_old_connections()
            try:
                barriere.wait(timeout=10)
                return fonction()
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=len(fonctions)) as pool:
            return list(pool.map(executer, fonctions))

    def test_deux_suppressions_un_seul_journal(self):
        ligne = ajouter_cafes(self.personne, 1, employe=self.barista)[0]
        resultats = self.simultanement([
            lambda: supprimer_clients([self.personne.pk], employe=self.gerant),
            lambda: supprimer_clients([self.personne.pk], employe=self.gerant),
        ])
        self.assertEqual(sorted(resultats), [0, 1])
        self.assertEqual(LogEntry.objects.filter(action_flag=DELETION).count(), 1)
        self.assertEqual(Client.objects.count(), 1)
        ligne.refresh_from_db()
        self.assertEqual((ligne.client_id, ligne.client_origine_id), (-1, self.personne.pk))

    def test_achat_concurrent_conserve_achat_ou_refuse_client_disparu(self):
        def acheter():
            try:
                return ajouter_cafes(self.personne, 1, employe=self.barista)
            except Client.DoesNotExist:
                return []

        _, lignes = self.simultanement([
            lambda: supprimer_clients([self.personne.pk], employe=self.gerant), acheter,
        ])
        self.assertFalse(Client.objects.filter(pk=self.personne.pk).exists())
        self.assertEqual(TransactionCafe.objects.count(), len(lignes))
        self.assertFalse(TransactionCafe.objects.exclude(client_id=-1).exists())
        self.assertEqual(Client.objects.get(pk=-1).nombre_cafes_achetes, 0)

    def test_annulation_concurrente_conserve_trace_et_fictif_a_zero(self):
        ligne = ajouter_cafes(self.personne, 1, employe=self.barista)[0]
        self.simultanement([
            lambda: supprimer_clients([self.personne.pk], employe=self.gerant),
            lambda: annuler_operation(ligne.pk, self.gerant),
        ])
        ligne.refresh_from_db()
        self.assertEqual((ligne.client_id, ligne.client_origine_id), (-1, self.personne.pk))
        self.assertTrue(ligne.annulee)
        self.assertEqual(AnnulationTransaction.objects.count(), 1)
        self.assertEqual(Client.objects.get(pk=-1).nombre_cafes_achetes, 0)

    def test_suppressions_clients_distincts_partagent_fictif(self):
        autre = Client.objects.create(nom_complet="Alex", telephone="418 555-2000")
        ajouter_cafes(self.personne, 1, employe=self.barista)
        ajouter_cafes(autre, 1, employe=self.barista)
        self.assertEqual(self.simultanement([
            lambda: supprimer_clients([self.personne.pk], employe=self.gerant),
            lambda: supprimer_clients([autre.pk], employe=self.gerant),
        ]), [1, 1])
        self.assertEqual(Client.objects.count(), 1)
        self.assertEqual(TransactionCafe.objects.filter(client_id=-1).count(), 2)
        self.assertEqual(LogEntry.objects.filter(action_flag=DELETION).count(), 2)


class SuppressionMigrationTests(TransactionTestCase):
    def test_migration_conserve_clients_transactions_et_annulations(self):
        avant = [("CafeDolbeau", "0012_parametres_gratuite")]
        apres = [("CafeDolbeau", "0013_transactioncafe_client_origine_id")]
        executor = MigrationExecutor(connection)
        dernier_schema = executor.loader.graph.leaf_nodes()
        try:
            executor.migrate(avant)
            apps = executor.loader.project_state(avant).apps
            AncienClient = apps.get_model("CafeDolbeau", "Client")
            AncienneTransaction = apps.get_model("CafeDolbeau", "TransactionCafe")
            AncienneAnnulation = apps.get_model("CafeDolbeau", "AnnulationTransaction")
            User = apps.get_model("auth", "User")
            personne = AncienClient.objects.create(nom_complet="Migration", telephone="418 555-3000")
            employe = User.objects.create(username="employe-migration")
            ligne = AncienneTransaction.objects.create(client=personne, quantite=3, employe_id=employe.pk)
            AncienneAnnulation.objects.create(transaction=ligne, employe_id=employe.pk)
            clients = list(AncienClient.objects.values())
            transactions = list(AncienneTransaction.objects.values())
            annulations = list(AncienneAnnulation.objects.values())
            executor = MigrationExecutor(connection)
            executor.migrate(apres)
            apps = executor.loader.project_state(apres).apps
            self.assertEqual(list(apps.get_model("CafeDolbeau", "Client").objects.values()), clients)
            self.assertEqual(list(apps.get_model("CafeDolbeau", "TransactionCafe").objects.values()), [
                {**ligne, "client_origine_id": None} for ligne in transactions
            ])
            self.assertEqual(list(apps.get_model("CafeDolbeau", "AnnulationTransaction").objects.values()), annulations)
            self.assertFalse(apps.get_model("CafeDolbeau", "Client").objects.filter(pk=-1).exists())
        finally:
            MigrationExecutor(connection).migrate(dernier_schema)
